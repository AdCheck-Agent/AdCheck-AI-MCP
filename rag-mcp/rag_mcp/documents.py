import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path

import yaml
from pydantic import ValidationError

from rag_mcp.config import Settings
from rag_mcp.errors import DocumentError
from rag_mcp.retrieval.embeddings import Embeddings
from rag_mcp.schemas import Chunk, DocumentRecord, PolicyRow

# 기업과 문서는 URL(샘플이면 파일 경로)로 정한다. 예전 파일에 이 두 컬럼이
# 남아 있을 수 있어서 받기는 하지만 값은 무시한다.
IDENTITY = ("company_id", "document_id")
COLUMNS = [name for name in PolicyRow.model_fields if name not in IDENTITY]
MARKDOWN_LINK = re.compile(r"\[([^\]]+)\]\([^)\s]+\)")


def _table(filename: str, data: bytes, settings: Settings) -> list[list]:
    suffix = Path(filename).suffix.lower()
    if len(data) > settings.rag_max_upload_bytes:
        raise DocumentError("파일 크기 제한을 초과했습니다.", "FILE_TOO_LARGE", 413)
    try:
        if suffix == ".csv":
            table = []
            for row in csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True):
                if len(table) > settings.rag_max_rows:
                    raise DocumentError("행 수 제한을 초과했습니다.")
                table.append(row)
            return table
        if suffix != ".xlsx":
            raise DocumentError(
                "CSV 또는 XLSX만 지원합니다. PDF는 지원하지 않습니다.", "UNSUPPORTED_FILE_TYPE", 415
            )
        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise DocumentError("XLSX 압축 해제 크기 제한을 초과했습니다.")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if wb.sheetnames != ["policies"]:
                raise DocumentError("XLSX는 policies 시트 하나만 포함해야 합니다.")
            sheet = wb["policies"]
            sheet.reset_dimensions()
            table = []
            for cells in sheet.iter_rows():
                if len(table) > settings.rag_max_rows:
                    raise DocumentError("행 수 제한을 초과했습니다.")
                if any(cell.data_type == "f" for cell in cells):
                    raise DocumentError("수식은 지원하지 않습니다. 계산된 값을 넣어 주세요.")
                table.append([cell.value for cell in cells])
            # 엑셀은 끝의 빈 칸을 생략하므로, 비어 있는 선택 컬럼이 짧게 들어온다.
            width = len(table[0]) if table else 0
            return [row + [None] * (width - len(row)) for row in table]
        finally:
            wb.close()
    except (UnicodeError, csv.Error, zipfile.BadZipFile, KeyError, ValueError) as exc:
        if isinstance(exc, DocumentError):
            raise
        raise DocumentError(
            "파일을 읽을 수 없습니다. UTF-8 CSV 또는 정상 XLSX인지 확인하세요."
        ) from exc


def sample_identity(path: Path) -> tuple[str, str]:
    """레포 샘플의 ID: data/corporate/{company_id}/{type}.csv -> company-type."""
    return path.parent.name, f"{path.parent.name}-{path.stem}"


def parse_policies(
    filename: str, data: bytes, settings: Settings, company_id: str, document_id: str
) -> list[PolicyRow]:
    """기업 파일 하나를 읽는다. company_id/document_id는 파일이 아니라 호출하는 쪽에서 받는다."""
    try:
        PolicyRow.model_validate(
            {"row_id": "x", "document_type": "corporate_values", "title": "x", "section": "x"}
            | {"content": "x", "company_id": company_id, "document_id": document_id}
        )
    except ValidationError as exc:
        raise DocumentError("기업·문서 ID 형식을 확인하세요.", "INVALID_IDENTITY") from exc
    table = _table(filename, data, settings)
    if not table:
        raise DocumentError("헤더와 최소 한 개의 정책 행이 필요합니다.")
    header = [str(x).strip() if x is not None else "" for x in table[0]]
    if len(set(header)) != len(header) or set(header) - set(IDENTITY) != set(COLUMNS):
        raise DocumentError(f"컬럼을 확인하세요. 필수 헤더: {', '.join(COLUMNS)}")
    rows = []
    seen = set()
    for number, values in enumerate(table[1:], 2):
        if all(x in (None, "") for x in values):
            continue
        if len(values) != len(header):
            raise DocumentError(f"{number}행: 컬럼 수가 헤더와 다릅니다.")
        obj = dict(zip(header, values, strict=True)) | {
            "company_id": company_id,
            "document_id": document_id,
        }
        for key in ("source_url", "published_at"):
            if obj[key] in (None, ""):
                obj[key] = None
        if isinstance(obj["published_at"], datetime):
            obj["published_at"] = obj["published_at"].date()
        try:
            row = PolicyRow.model_validate(obj)
        except ValidationError as exc:
            details = "; ".join(f"{e['loc'][0]}: {e['msg']}" for e in exc.errors())
            raise DocumentError(f"{number}행: {details}") from exc
        if row.row_id in seen:
            raise DocumentError(f"{number}행: 중복 row_id입니다.")
        seen.add(row.row_id)
        rows.append(row)
    if not rows:
        raise DocumentError("최소 한 개의 정책 행이 필요합니다.")
    if any(
        any(getattr(r, key) != getattr(rows[0], key) for key in ("document_type", "title"))
        for r in rows
    ):
        raise DocumentError("한 파일에는 같은 문서 유형·제목만 넣으세요.")
    return rows


async def corporate_document(rows: list[PolicyRow], embeddings: Embeddings) -> DocumentRecord:
    first = rows[0]
    chunks = []
    for row in rows:
        for i, content in enumerate(await embeddings.split(row.content)):
            chunks.append(
                Chunk(
                    chunk_id=f"{row.row_id}:{i}",
                    section=row.section,
                    content=content,
                    metadata={
                        "row_id": row.row_id,
                        # 파일 순서. 첨부 문서 검색은 조각을 이 순서로 돌려준다.
                        "ordinal": len(chunks),
                        "source_urls": [str(row.source_url)] if row.source_url else [],
                        "published_at": row.published_at.isoformat() if row.published_at else None,
                    },
                )
            )
    canonical = json.dumps([r.model_dump(mode="json") for r in rows], sort_keys=True)
    return DocumentRecord(
        company_id=first.company_id,
        document_id=first.document_id,
        corpus="corporate",
        document_type=first.document_type,
        title=first.title,
        checksum=hashlib.sha256(canonical.encode()).hexdigest(),
        chunks=chunks,
        metadata={
            "category": "corporate_value",
            "verification_status": "user_provided_not_independently_verified",
            "record_type": "uploaded_company_document",
        },
    )


async def social_documents(data_dir: Path, embeddings: Embeddings) -> list[DocumentRecord]:
    manifest = json.loads((data_dir / "manifest.json").read_text("utf-8"))
    registry = json.loads((data_dir / "sources.json").read_text("utf-8"))
    sources = {s["source_id"]: s["url"] for s in registry["sources"]}
    documents = []
    for item in manifest["documents"]:
        if item["corpus"] != "social":
            continue
        path = (data_dir / item["path"]).resolve()
        if not path.is_relative_to((data_dir / "social").resolve()):
            raise DocumentError("사회 문서 manifest 경로가 social 범위를 벗어났습니다.")
        text = path.read_text("utf-8")
        _, front, body = text.split("---", 2)
        meta = yaml.safe_load(front)
        if meta["document_id"] != item["document_id"] or not meta["retrieval_enabled"]:
            raise DocumentError("사회 문서 manifest와 metadata가 일치하지 않습니다.")
        chunks, section = [], meta["title"]
        # 사실과 프로젝트 안내가 섞이지 않게 제목마다 자기 출처를 유지한다.
        sections = re.split(r"(?m)^#{1,6} (.+)\n", body)
        for i, part in enumerate(sections):
            if i % 2:
                section = part.strip()
                continue
            if not part.strip():
                continue
            # 링크 주소는 source_urls에 이미 있다(또는 레포 안 파일을 가리킨다). 이를 임베딩하면
            # URL만 든 조각이 생겨 무관한 질문에도 상위에 올랐다.
            for content in await embeddings.split(MARKDOWN_LINK.sub(r"\1", part).strip()):
                chunks.append(
                    Chunk(chunk_id=f"section:{len(chunks)}", section=section, content=content)
                )
        documents.append(
            DocumentRecord(
                company_id=None,
                document_id=meta["document_id"],
                corpus="social",
                document_type=meta["record_type"],
                title=meta["title"],
                checksum=hashlib.sha256(text.encode()).hexdigest(),
                chunks=chunks,
                metadata={
                    "rule_id": meta["rule_id"],
                    "record_type": meta["record_type"],
                    "category": meta["category"],
                    "verification_status": meta["verification_status"],
                    "source_ids": meta["source_ids"],
                    "source_urls": [sources[s] for s in meta["source_ids"]],
                },
            )
        )
    return documents
